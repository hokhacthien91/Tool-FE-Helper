"""test_build.py — regression test cho PPTX Master Builder.

Chạy:  python3 test_build.py

Build test_deck.pptx (+ các file thật nếu có) rồi kiểm:
  1. Build không lỗi
  2. verify.py: id unique, ảnh resolved, không dup placeholder
  3. Số slide output == input (build không mất/thêm slide)
  4. sp_children khớp rep.shapes mọi slide (bug #1 — cxnSp lệch zip)
  5. Slide non-rep giữ layout gốc (bug A — không leak content rep)
  6. Visual verify: mọi slide render được >= ngưỡng (nếu có renderer)

Thoát code 0 nếu PASS, 1 nếu có FAIL.
"""
import os
import sys
import tempfile

from pptx import Presentation
from build import build, SHAPE_TAGS
from classify import analyze
import verify
import preview

HERE = os.path.dirname(os.path.abspath(__file__))
PARENT = os.path.dirname(HERE)

# file test: test_deck.pptx (bắt buộc) + file thật nếu tìm thấy
TEST_FILES = [os.path.join(HERE, "test_deck.pptx")]
for name in ("Sales Deck Template.pptx", "DC Sales Deck_v2.pptx"):
    p = os.path.join(PARENT, name)
    if os.path.isfile(p):
        TEST_FILES.append(p)


def _check(cond, msg, fails):
    print(("  ✅ " if cond else "  ❌ ") + msg)
    if not cond:
        fails.append(msg)


def test_file(src, do_visual=True):
    print(f"\n=== {os.path.basename(src)} ===")
    fails = []
    n_in = len(Presentation(src).slides._sldIdLst)

    # 4. sp_children khớp rep.shapes (bug #1)
    from pptx.oxml.ns import qn  # noqa
    prs_in = Presentation(src)
    slides_info = analyze(src)
    misaligned = []
    for i, sl in enumerate(prs_in.slides):
        spc = len([c for c in sl.shapes._spTree if c.tag in SHAPE_TAGS])
        walk = len(slides_info[i].shapes)
        if spc != walk:
            misaligned.append((i + 1, spc, walk))
    _check(not misaligned, f"shape↔role aligned (0 misaligned), got {misaligned[:3]}", fails)

    # 1. build không lỗi
    out = os.path.join(tempfile.gettempdir(), "rt_" + os.path.basename(src))
    try:
        created, renamed = build(src, out, verbose=False, names={}, options={})
        _check(True, f"build OK ({len(created)} layouts)", fails)
    except Exception as e:
        _check(False, f"build raised: {e}", fails)
        return fails

    # 3. số slide giữ nguyên
    n_out = len(Presentation(out).slides._sldIdLst)
    _check(n_out == n_in, f"slide count preserved ({n_in} -> {n_out})", fails)

    # 2. structure verify
    ok = verify.verify(out)
    _check(ok, "structure verify (id/img/ph)", fails)

    # 5. slide non-rep giữ layout gốc (bug A) — so layout name src vs out cho slide KHÔNG phải rep
    ps, po = Presentation(src), Presentation(out)
    from classify import group_slides
    groups = group_slides(analyze(src))
    reps = {members[0].index for members in groups.values()}
    leaked = []
    src_lay = [s.slide_layout.name for s in ps.slides]
    out_lay = [s.slide_layout.name for s in po.slides]
    for i in range(n_in):
        if (i + 1) not in reps and src_lay[i] != out_lay[i]:
            leaked.append((i + 1, src_lay[i], out_lay[i]))
    _check(not leaked, f"non-rep slides keep layout (no leak), got {leaked[:3]}", fails)

    # 6. visual verify (nếu có renderer)
    if do_visual and preview.renderer_name():
        preview._CACHE.clear()
        r = preview.verify_visual(src, out)
        bad = [(s["index"], s["pct"]) for s in r["slides"]
               if s["status"] == "diff"]
        _check(not bad, f"visual: no slide differs (<97%), got {bad[:3]}", fails)
    return fails


def main():
    if not os.path.isfile(TEST_FILES[0]):
        print("⚠ Chưa có test_deck.pptx — chạy: python3 make_test_deck.py")
        sys.exit(1)
    all_fails = []
    for f in TEST_FILES:
        # visual verify chỉ chạy test_deck (nhỏ, nhanh); file V2 lớn -> bỏ visual cho nhanh
        all_fails += test_file(f, do_visual=("test_deck" in f))
    print("\n" + "=" * 40)
    if all_fails:
        print(f"❌ FAIL — {len(all_fails)} check failed")
        sys.exit(1)
    print("✅ ALL PASS")
    sys.exit(0)


if __name__ == "__main__":
    main()
